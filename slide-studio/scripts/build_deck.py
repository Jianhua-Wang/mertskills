# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0", "pyyaml>=6", "pillow>=10"]
# ///
"""Build a .pptx from a deck.yaml spec on the slide-studio academic template.

Usage:
    uv run scripts/build_deck.py deck.yaml [-o deck.pptx] [--template PATH]
    uv run scripts/build_deck.py --layouts          # list layouts and their fields

Every slide is placed on one of the template's layouts and every piece of text
goes into a layout placeholder, so position, size, font and colour always come
from the template. After building, the deck is checked: text that overflows its
slot, headlines over two lines or naming a topic instead of a finding, wordy
points and slides, figures with low resolution or the wrong aspect ratio,
unrequested speaker notes, and talk-length problems are reported. Errors make
the exit code 1.

See references/deck-spec.md for the YAML format.
"""

from __future__ import annotations

import argparse
import copy
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import yaml
from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import design as D  # noqa: E402

SKILL_DIR = Path(__file__).resolve().parent.parent
DEFAULT_TEMPLATE = SKILL_DIR / "assets" / "academic-template.pptx"

# Roles that take a list of points (nested lists are sub-points).
LIST_ROLES = {"text", "left_text", "right_text", "points"}
# Fields every slide may carry besides its layout's roles.
COMMON_KEYS = {"layout", "notes", "reveal", "slide_number", "backup", "highlight", "table", "left",
               "right", "id", "borrowed"}
CONTENT_LAYOUTS = {"statement", "figure", "figure_text", "two_figures", "text", "two_columns",
                   "table", "summary"}
CJK_RE = re.compile(r"[\u2e80-\u9fff\uac00-\ud7af\uf900-\ufaff\uff00-\uffef\u3000-\u303f]")
# Conciseness budgets, in words (two CJK characters count as one word).
MAX_POINT_WORDS = 12
MAX_SLIDE_WORDS = 40
# Headlines that name a topic ("QC results", "Distribution of RIN") instead of
# stating a finding ("All 48 samples pass QC").
_TOPIC = (r"(analysis|distributions?|comparisons?|evaluation|overview|summary|results?|statistics|"
          r"characteri[sz]ation|effects?|impact|performance|profiles?|assessment|landscape)")
LABEL_START_RE = re.compile(rf"^(an? |the )?{_TOPIC} (of|between|across|for|on|in)\b", re.I)
LABEL_PREFIX_RE = re.compile(r"^(results?|methods?|background|motivation|introduction|overview|summary|"
                             r"setup|qc|quality control|conclusions?|discussion|outline|"
                             r"结果|方法|背景|介绍|概述|总结|讨论|结论)\s*[:：—–]", re.I)
LABEL_END_RE = re.compile(r"\b(results?|overview|distributions?|comparison|analysis|statistics|summary)$", re.I)
LABEL_END_CJK = ("结果", "分析", "概述", "概况", "分布", "比较", "对比", "介绍", "背景", "方法",
                 "总结", "流程", "统计", "情况", "展示", "示意图")
VS_RE = re.compile(r"\b(vs\.?|versus)(\s|$)", re.I)
NUMERIC_RE = re.compile(r"^[\s+\-\u2212\u00b1\u2013~<>=\u2264\u2265$%\u00d7.,0-9()\[\]*a-zA-Z/]*[0-9][\s0-9.,%*\u2020\u2021()]*[a-zA-Z%]*$")


# ==========================================================================
# Reporting
# ==========================================================================

@dataclass
class Issue:
    level: str   # ERROR | WARN
    where: str
    msg: str

    def __str__(self) -> str:
        return f"{self.level:5s} {self.where}: {self.msg}"


@dataclass
class Report:
    issues: list[Issue] = field(default_factory=list)

    def error(self, where: str, msg: str) -> None:
        self.issues.append(Issue("ERROR", where, msg))

    def warn(self, where: str, msg: str) -> None:
        self.issues.append(Issue("WARN", where, msg))

    @property
    def errors(self) -> list[Issue]:
        return [i for i in self.issues if i.level == "ERROR"]


class SpecError(Exception):
    pass


# ==========================================================================
# Inline markup:  **bold**  *italic*  ==accent==  ^{sup}  _{sub}
# ==========================================================================

TOKEN_RE = re.compile(
    r"\*\*(?P<b>.+?)\*\*"
    r"|==(?P<acc>.+?)=="
    r"|(?<![\w*])\*(?P<i>[^*\s][^*]*?)\*(?![\w*])"
    r"|\^\{(?P<sup>[^}]*)\}"
    r"|_\{(?P<sub>[^}]*)\}"
)


@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    accent: bool = False
    baseline: int = 0   # percent; +30 superscript, -25 subscript


def parse_markup(text: str) -> list[Run]:
    runs: list[Run] = []
    pos = 0
    for m in TOKEN_RE.finditer(text):
        if m.start() > pos:
            runs.append(Run(text[pos:m.start()]))
        if m.group("b") is not None:
            runs.append(Run(m.group("b"), bold=True))
        elif m.group("acc") is not None:
            runs.append(Run(m.group("acc"), bold=True, accent=True))
        elif m.group("i") is not None:
            runs.append(Run(m.group("i"), italic=True))
        elif m.group("sup") is not None:
            runs.append(Run(m.group("sup"), baseline=30))
        else:
            runs.append(Run(m.group("sub"), baseline=-25))
        pos = m.end()
    if pos < len(text):
        runs.append(Run(text[pos:]))
    return [r for r in runs if r.text]


def plain(text: str) -> str:
    return "".join(r.text for r in parse_markup(text))


# ==========================================================================
# Text measurement
# ==========================================================================

FONT_CANDIDATES = {
    "regular": [
        "/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts/arial.ttf",
        "/Library/Fonts/Arial.ttf",
        "/System/Library/Fonts/Supplementary/Arial.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "/usr/share/fonts/truetype/msttcorefonts/Arial.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    ],
    "bold": [
        "/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts/arialbd.ttf",
        "/Library/Fonts/Arial Bold.ttf",
        "/System/Library/Fonts/Supplementary/Arial Bold.ttf",
        "C:/Windows/Fonts/arialbd.ttf",
        "/usr/share/fonts/truetype/msttcorefonts/Arial_Bold.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
    ],
    "cjk": [
        "/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts/Deng.ttf",
        "C:/Windows/Fonts/Deng.ttf",
    ],
    "cjk_bold": [
        "/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts/Dengb.ttf",
        "C:/Windows/Fonts/Dengb.ttf",
    ],
}
SCALE = 20  # measure at 20 px per pt for sub-point precision


class Measurer:
    """Estimate how many lines a paragraph wraps to in PowerPoint."""

    def __init__(self) -> None:
        self._fonts: dict[tuple[str, int], ImageFont.FreeTypeFont | None] = {}
        self.missing: set[str] = set()

    def _font(self, kind: str, size_pt: float):
        key = (kind, int(size_pt * SCALE))
        if key not in self._fonts:
            font = None
            for path in FONT_CANDIDATES[kind]:
                if Path(path).exists():
                    font = ImageFont.truetype(path, key[1])
                    break
            if font is None:
                self.missing.add(kind)
            self._fonts[key] = font
        return self._fonts[key]

    def width_pt(self, text: str, size_pt: float, bold: bool) -> float:
        if not text:
            return 0.0
        total = 0.0
        # split into CJK and non-CJK stretches
        for chunk in re.findall(rf"{CJK_RE.pattern}+|[^\u2e80-\u9fff\uac00-\ud7af\uf900-\ufaff\uff00-\uffef\u3000-\u303f]+", text):
            cjk = bool(CJK_RE.match(chunk))
            kind = ("cjk_bold" if bold else "cjk") if cjk else ("bold" if bold else "regular")
            font = self._font(kind, size_pt)
            if font is None:
                total += len(chunk) * size_pt * (1.0 if cjk else 0.56)
            else:
                total += font.getlength(chunk) / SCALE
        return total

    def lines(self, runs: list[Run], size_pt: float, bold: bool, width_in: float) -> int:
        """Greedy word wrap; CJK characters break anywhere."""
        limit = width_in * 72
        tokens: list[tuple[str, bool]] = []
        for r in runs:
            b = bold or r.bold
            for tok in re.findall(rf"{CJK_RE.pattern}|\s+|[^\s\u2e80-\u9fff\uac00-\ud7af\uf900-\ufaff\uff00-\uffef\u3000-\u303f]+", r.text):
                tokens.append((tok, b))
        n, cur = 1, 0.0
        for tok, b in tokens:
            w = self.width_pt(tok, size_pt, b)
            if tok.isspace():
                cur += w
                continue
            if cur + w > limit and cur > 0:
                n += 1
                cur = w
            else:
                cur += w
        return n


MEASURE = Measurer()


def text_height_in(paras: list[tuple[int, list[Run]]], style: D.TextStyle, width_in: float) -> tuple[float, list[int]]:
    """Height of a list of (level, runs) paragraphs in a slot of the given width."""
    h_pt = 0.0
    per_para: list[int] = []
    for i, (lvl, runs) in enumerate(paras):
        size = style.sizes[min(lvl, len(style.sizes) - 1)]
        if style.bullets == "dot":
            indent = D.BULLET_INDENT * (lvl + 1)
        elif style.bullets == "number":
            indent = 0.5
        else:
            indent = 0.0
        n = MEASURE.lines(runs, size, style.bold, width_in - indent)
        per_para.append(n)
        if i > 0:
            h_pt += style.space_before[min(lvl, len(style.space_before) - 1)]
        h_pt += n * size * D.LINE_HEIGHT * style.line_spacing / 100
    return h_pt / 72, per_para


# ==========================================================================
# Writing text into placeholders
# ==========================================================================

def flatten_items(items, level: int = 0) -> list[tuple[int, str]]:
    """YAML list -> [(level, text)]. A nested list holds the previous item's sub-points."""
    if isinstance(items, str):
        return [(level, line) for line in items.split("\n") if line.strip()]
    out: list[tuple[int, str]] = []
    for it in items or []:
        if isinstance(it, list):
            out.extend(flatten_items(it, level + 1))
        elif isinstance(it, dict) and len(it) == 1:
            # "- Point: [sub, sub]" written as a YAML mapping
            (k, v), = it.items()
            out.append((level, str(k)))
            out.extend(flatten_items(v, level + 1))
        else:
            out.append((level, str(it)))
    return out


def set_runs(paragraph, runs: list[Run]) -> None:
    for r in runs:
        run = paragraph.add_run()
        run.text = r.text
        f = run.font
        if r.bold:
            f.bold = True
        if r.italic:
            f.italic = True
        if r.accent:
            f.color.theme_color = MSO_THEME_COLOR.ACCENT_1
        rpr = run._r.get_or_add_rPr()
        if r.baseline:
            rpr.set("baseline", str(r.baseline * 1000))
        if CJK_RE.search(r.text):
            rpr.set("lang", "zh-CN")
            rpr.set("altLang", "en-US")
        else:
            rpr.set("lang", "en-US")


def fill_text(ph, paras: list[tuple[int, str]]) -> None:
    tf = ph.text_frame
    # drop any paragraphs beyond the first, then clear it
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    first = tf.paragraphs[0]
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    for i, (lvl, text) in enumerate(paras):
        p = first if i == 0 else tf.add_paragraph()
        if lvl:
            p.level = lvl
        set_runs(p, parse_markup(text))


# ==========================================================================
# Deck builder
# ==========================================================================

@dataclass
class Ctx:
    base: Path
    meta: dict
    report: Report
    section_no: int = 0
    content_slides: int = 0
    slides_without_notes: list[str] = field(default_factory=list)
    slides_with_notes: list[str] = field(default_factory=list)


def placeholder_for(slide, slot: D.Slot):
    idx = 0 if slot.idx is None else slot.idx
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    return None


def remove_shape(shape) -> None:
    el = shape._element
    el.getparent().remove(el)


def resolve(ctx: Ctx, rel: str) -> Path:
    p = Path(rel).expanduser()
    return p if p.is_absolute() else (ctx.base / p)


def word_count(text: str) -> float:
    """Words in a text, with two CJK characters counted as one word."""
    s = plain(text)
    cjk = len(CJK_RE.findall(s))
    latin = [w for w in CJK_RE.sub(" ", s).split() if any(c.isalnum() for c in w)]
    return len(latin) + cjk / 2


def is_label_headline(text: str) -> bool:
    """True when a headline names a topic ("QC results") instead of stating a finding."""
    s = plain(text).strip().rstrip(".。!！?？")
    words = s.split()
    if not CJK_RE.search(s) and len(words) < 4:
        return True
    if LABEL_PREFIX_RE.search(s) or LABEL_START_RE.search(s) or LABEL_END_RE.search(s):
        return True
    if s.endswith(LABEL_END_CJK):
        return True
    return bool(VS_RE.search(s)) and word_count(s) < 7


def check_text(ctx: Ctx, where: str, slot: D.Slot, paras: list[tuple[int, str]],
               claim: bool = True) -> float:
    """Lint one text slot and return the measured text height in inches.

    claim=False skips the finding-not-topic headline check (summary, backup slides)."""
    style = D.TEXT_STYLES[slot.style]
    runs = [(lvl, parse_markup(t)) for lvl, t in paras]
    h, lines = text_height_in(runs, style, slot.w)
    if h > slot.limit_h + 0.02:
        ctx.report.error(where, f"{slot.role} overflows its slot ({h:.2f} in of text in "
                                f"{slot.limit_h:.2f} in); cut words, not font size")
    if slot.role == "headline":
        if lines and lines[0] > 2:
            ctx.report.error(where, f"headline wraps to {lines[0]} lines (max 2)")
        if claim and paras and is_label_headline(paras[0][1]):
            ctx.report.warn(where, f"headline names a topic, not a finding: '{plain(paras[0][1])[:50]}'; "
                                   "say what the slide shows (e.g. 'All 48 samples pass QC', "
                                   "'X beats Y on 5 of 6 datasets')")
        if paras and paras[0][1].rstrip().endswith(":"):
            ctx.report.warn(where, "headline ends with ':'; state the finding instead")
    if slot.role in LIST_ROLES:
        top = [t for lvl, t in paras if lvl == 0]
        limit = 4 if slot.role == "points" else 5
        if len(top) > limit:
            ctx.report.warn(where, f"{len(top)} top-level points in {slot.role} (keep to {limit} or fewer)")
        for lvl, t in paras:
            if lvl > 1:
                ctx.report.warn(where, "three levels of bullets; flatten to two")
                break
        for _, t in paras:
            n = word_count(t)
            if n > MAX_POINT_WORDS:
                ctx.report.warn(where, f"point has {n:.0f} words (keep to {MAX_POINT_WORDS}): "
                                       f"'{plain(t)[:40]}...'")
        if any(plain(t).lstrip().startswith(("\u2022", "- ", "* ")) for _, t in paras):
            ctx.report.warn(where, "literal bullet character in text; bullets are automatic")
    return h


def place_picture(ctx: Ctx, slide, slot: D.Slot, image: str, where: str, borrowed: bool = False):
    path = resolve(ctx, image)
    if not path.exists():
        ctx.report.error(where, f"image not found: {path}")
        return None
    if path.suffix.lower() not in (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif"):
        ctx.report.error(where, f"{path.name}: use PNG (300 dpi) or JPEG; convert PDF/SVG first")
        return None
    with Image.open(path) as im:
        px_w, px_h = im.size
    ph = placeholder_for(slide, slot)
    # Fit the whole image inside the slot (never crop), centred.
    scale = min(slot.w / px_w, slot.h / px_h)
    w, h = px_w * scale, px_h * scale
    x = slot.x + (slot.w - w) / 2
    y = slot.y + (slot.h - h) / 2
    if ph is not None:
        pic = ph.insert_picture(str(path))
        pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0
        pic.left, pic.top, pic.width, pic.height = Inches(x), Inches(y), Inches(w), Inches(h)
    else:
        pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    dpi = px_w / w
    if dpi < 150:
        ctx.report.warn(where, f"{path.name} is {dpi:.0f} dpi at display size (export at 300 dpi)")
    fill = (w * h) / (slot.w * slot.h)
    if fill < 0.85 and not borrowed:   # a figure taken from a paper cannot be re-exported
        ctx.report.warn(where, f"{path.name} fills {fill:.0%} of its slot; export it at "
                               f"{slot.w:.2f} x {slot.h:.2f} in for a full fit")
    return (x, y, w, h)


def draw_highlights(ctx: Ctx, slide, box_in, highlights, where: str) -> None:
    if not highlights or box_in is None:
        return
    ix, iy, iw, ih = box_in
    for hl in highlights:
        try:
            x0, y0, x1, y1 = [float(v) for v in hl["box"]]
        except (KeyError, TypeError, ValueError):
            ctx.report.error(where, "highlight needs box: [x0, y0, x1, y1] as fractions of the image")
            continue
        bx, by = ix + x0 * iw, iy + y0 * ih
        bw, bh = (x1 - x0) * iw, (y1 - y0) * ih
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(bx), Inches(by), Inches(bw), Inches(bh))
        shp.name = "Highlight"
        shp.fill.background()
        shp.line.color.theme_color = MSO_THEME_COLOR.ACCENT_1
        shp.line.width = Pt(2.25)
        shp.shadow.inherit = False
        label = hl.get("label")
        if label:
            lh = 0.32
            ly = by - lh - 0.04 if by - lh - 0.04 >= D.BODY_Y - 0.1 else by + bh + 0.04
            tb = slide.shapes.add_textbox(Inches(bx), Inches(ly), Inches(max(bw, 2.5)), Inches(lh))
            tb.name = "Highlight label"
            tf = tb.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.word_wrap = False
            tf.vertical_anchor = MSO_ANCHOR.BOTTOM if ly < by else MSO_ANCHOR.TOP
            p = tf.paragraphs[0]
            set_runs(p, [Run(r.text, bold=True, accent=True, italic=r.italic, baseline=r.baseline)
                         for r in parse_markup(str(label))])
            for run in p.runs:
                run.font.size = Pt(14)


def set_cell_border(cell, top: float | None, bottom: float | None) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB", "a:solidFill", "a:noFill"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)

    def line(tag: str, width_pt: float | None):
        el = etree.SubElement(tcPr, qn(tag))
        if width_pt:
            el.set("w", str(int(width_pt * D.EMU_PER_PT)))
            el.set("cap", "flat")
            el.set("cmpd", "sng")
            fill = etree.SubElement(el, qn("a:solidFill"))
            etree.SubElement(fill, qn("a:srgbClr")).set("val", D.INK)
            etree.SubElement(el, qn("a:prstDash")).set("val", "solid")
        else:
            el.set("w", "0")
            etree.SubElement(el, qn("a:noFill"))

    line("a:lnL", None)
    line("a:lnR", None)
    line("a:lnT", top)
    line("a:lnB", bottom)


def set_cell_fill(cell, hex_color: str | None) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    if hex_color:
        fill = etree.SubElement(tcPr, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", hex_color)
    else:
        etree.SubElement(tcPr, qn("a:noFill"))


def draw_table(ctx: Ctx, slide, spec: dict, where: str) -> None:
    header = [str(c) for c in spec.get("header") or spec.get("columns") or []]
    rows = [[("" if c is None else str(c)) for c in r] for r in spec.get("rows") or []]
    if not header or not rows:
        ctx.report.error(where, "table needs header: [...] and rows: [[...], ...]")
        return
    ncol = len(header)
    if any(len(r) != ncol for r in rows):
        ctx.report.error(where, f"every table row needs {ncol} cells")
        return
    if len(rows) + 1 > D.TABLE_MAX_ROWS:
        ctx.report.error(where, f"{len(rows)} body rows (max {D.TABLE_MAX_ROWS - 1}); "
                                "keep the rows that carry the message, move the rest to backup")
    _, y0, region_w, region_h = D.TABLE_REGION
    width = float(spec.get("width") or min(region_w, max(7.5, 2.4 * ncol)))
    width = min(width, region_w)
    x0 = (D.SLIDE_W - width) / 2   # tables sit centred under the headline
    weights = [float(w) for w in (spec.get("widths") or [1] * ncol)]
    col_w = [width * w / sum(weights) for w in weights]
    aligns = spec.get("align")
    if not aligns:
        aligns = ["l"] + [
            "r" if all(NUMERIC_RE.match(plain(r[j]).strip() or "0") for r in rows) else "l"
            for j in range(1, ncol)
        ]
    highlight = {int(i) for i in (spec.get("highlight") or [])}

    nrow = len(rows) + 1
    gf = slide.shapes.add_table(nrow, ncol, Inches(x0), Inches(y0), Inches(width),
                                Inches(D.TABLE_ROW_H * nrow))
    gf.name = "Table"
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for attr in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        if tblPr.get(attr) is not None:
            del tblPr.attrib[attr]
    style_id = tblPr.find(qn("a:tableStyleId"))
    if style_id is None:
        style_id = etree.SubElement(tblPr, qn("a:tableStyleId"))
    style_id.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for j, w in enumerate(col_w):
        tbl.columns[j].width = Inches(w)

    rule_top, rule_mid, rule_bot = D.TABLE_RULE_PT
    hs, bs = D.TEXT_STYLES["table_header"], D.TEXT_STYLES["table_body"]
    total_h = 0.0
    for i in range(nrow):
        tbl.rows[i].height = Inches(D.TABLE_ROW_H)
        cells = header if i == 0 else rows[i - 1]
        hl = i in highlight
        style = hs if i == 0 else bs
        row_lines = 1
        for j in range(ncol):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            set_cell_border(cell, rule_top if i == 0 else None,
                            rule_mid if i == 0 else (rule_bot if i == nrow - 1 else None))
            set_cell_fill(cell, D.ACCENT_TINT if hl else None)
            tf = cell.text_frame
            p = tf.paragraphs[0]
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[
                str(aligns[j])[0].lower()]
            runs = parse_markup(cells[j])
            set_runs(p, runs)
            for run in p.runs:
                run.font.size = Pt(style.sizes[0])
                if i == 0 or hl:
                    run.font.bold = True
                if run.font.color.type is None:
                    run.font.color.rgb = RGBColor.from_string(D.INK)
            n = MEASURE.lines(runs, style.sizes[0], i == 0 or hl, col_w[j] - 0.2)
            if n > 1:
                ctx.report.warn(where, f"table cell wraps: '{plain(cells[j])[:30]}' "
                                       "(shorten it or widen the column)")
            row_lines = max(row_lines, n)
        total_h += max(D.TABLE_ROW_H, row_lines * style.sizes[0] * D.LINE_HEIGHT / 72 + 0.08)
    if total_h > region_h:
        ctx.report.error(where, f"table is {total_h:.2f} in tall (max {region_h:.2f} in)")


def add_slide_number(slide, slot: D.Slot) -> None:
    """python-pptx does not copy the slide-number placeholder; add it."""
    tree = slide.shapes._spTree
    next_id = max([int(el.get("id")) for el in tree.iter(qn("p:cNvPr"))] + [1]) + 1
    xml = (
        '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<p:nvSpPr><p:cNvPr id="{next_id}" name="Slide Number"/>'
        '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
        f'<p:nvPr><p:ph type="sldNum" sz="quarter" idx="{slot.idx}"/></p:nvPr></p:nvSpPr>'
        '<p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/>'
        '<a:p><a:fld id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}" type="slidenum">'
        '<a:rPr lang="en-US"/><a:t>&lt;#&gt;</a:t></a:fld></a:p></p:txBody></p:sp>'
    )
    tree.append(etree.fromstring(xml))


def slide_fields(layout: D.Layout, spec: dict) -> dict[str, object]:
    """Map YAML keys (including left/right groups) to slot roles."""
    vals: dict[str, object] = {}
    for k, v in spec.items():
        if k in ("left", "right") and isinstance(v, dict):
            for sub, sv in v.items():
                if sub == "highlight":
                    continue
                role = {"label": "label", "heading": "heading", "image": "image",
                        "text": "text"}.get(sub, sub)
                vals[f"{k}_{role}"] = sv
        elif k not in COMMON_KEYS:
            vals[k] = v
    # "headline" is accepted as an alias for "title" on hero slides and vice versa
    if layout.slot("headline") is None and "headline" in vals and "title" not in vals:
        vals["title"] = vals.pop("headline")
    if layout.slot("headline") is not None and "title" in vals and "headline" not in vals:
        vals["headline"] = vals.pop("title")
    if layout.key == "statement" and "statement" in vals:
        vals["text"] = vals.pop("statement")
    return vals


def pack_stack(slide, layout: D.Layout, heights: dict[str, float]) -> None:
    """Stack the filled hero slots top to bottom and centre the block (see design.Stack)."""
    st = layout.stack
    items = []
    for role, gap in zip(st.roles, st.gaps):
        if role in heights:
            items.append((layout.slot(role), gap if items else 0.0, heights[role] + 0.04))
    if not items:
        return
    total = sum(gap + h for _, gap, h in items)
    y = max(st.top, st.top + (st.bottom - st.top - total) / 2 - st.lift)
    for slot, gap, h in items:
        y += gap
        ph = placeholder_for(slide, slot)
        # set all four so the slide gets a complete xfrm of its own
        ph.left, ph.top, ph.width, ph.height = Inches(slot.x), Inches(y), Inches(slot.w), Inches(h)
        y += h


def build_slide(prs, ctx: Ctx, spec: dict, n: int, *, visible: int | None = None,
                images: dict[str, str] | None = None) -> None:
    key = spec.get("layout")
    layout = D.LAYOUT_BY_KEY.get(key)
    if layout is None:
        raise SpecError(f"slide {n}: unknown layout {key!r}; choose from "
                        + ", ".join(D.LAYOUT_BY_KEY))
    where = f"slide {n} ({key})"
    vals = slide_fields(layout, spec)
    if images:
        vals.update(images)
    unknown = [k for k in vals if layout.slot(k) is None]
    if unknown:
        raise SpecError(f"{where}: fields {unknown} do not exist on this layout; it takes "
                        + ", ".join(s.role for s in layout.slots if s.kind != "sldNum")
                        + (", table" if key == "table" else ""))
    if key == "section" and "number" not in vals:
        ctx.section_no += 1
        vals["number"] = f"{ctx.section_no:02d}"
    elif key == "section":
        ctx.section_no += 1

    idx = D.LAYOUTS.index(layout)
    slide = prs.slides.add_slide(prs.slide_layouts[idx])
    pic_boxes: dict[str, tuple] = {}
    heights: dict[str, float] = {}
    body_words = 0.0
    claim = key != "summary" and not spec.get("backup")

    for slot in layout.slots:
        if slot.kind == "sldNum":
            continue
        ph = placeholder_for(slide, slot)
        val = vals.get(slot.role)
        if val is None or val == "" or val == []:
            if ph is not None:
                remove_shape(ph)
            if slot.role == "headline":
                ctx.report.error(where, "missing headline")
            continue
        if slot.kind == "pic":
            box = place_picture(ctx, slide, slot, str(val), where, bool(spec.get("borrowed")))
            if box:
                pic_boxes[slot.role] = box
            continue
        if slot.role in LIST_ROLES:
            paras = flatten_items(val)
        else:
            paras = [(0, line) for line in str(val).split("\n") if line.strip()]
        heights[slot.role] = check_text(ctx, where, slot, paras, claim)
        if slot.role in LIST_ROLES:
            body_words += sum(word_count(t) for _, t in paras)
        if visible is not None and slot.role in LIST_ROLES:
            shown, count = [], 0
            for lvl, t in paras:
                if lvl == 0:
                    count += 1
                if count <= visible:
                    shown.append((lvl, t))
            paras = shown
        fill_text(ph, paras)

    if body_words > MAX_SLIDE_WORDS:
        ctx.report.warn(where, f"{body_words:.0f} words of body text (keep to {MAX_SLIDE_WORDS}); "
                               "cut words or show it as a figure")
    if layout.stack is not None:
        pack_stack(slide, layout, heights)
    if spec.get("borrowed") and layout.slot("source") is not None and not vals.get("source"):
        ctx.report.error(where, "borrowed figure without a source; cite it (Author et al., Journal year)")

    # highlights: on single-image layouts under "highlight", on two_figures under left/right
    draw_highlights(ctx, slide, pic_boxes.get("image"), spec.get("highlight"), where)
    for side in ("left", "right"):
        if isinstance(spec.get(side), dict):
            draw_highlights(ctx, slide, pic_boxes.get(f"{side}_image"),
                            spec[side].get("highlight"), where)

    if key == "table":
        if "table" in spec:
            draw_table(ctx, slide, spec["table"], where)
        else:
            ctx.report.warn(where, "table layout without a table: draw your own content on it")

    num_slot = layout.slot("slide_number")
    if num_slot is not None and spec.get("slide_number", True) is not False:
        add_slide_number(slide, num_slot)

    notes = spec.get("notes")
    if notes:
        slide.notes_slide.notes_text_frame.text = str(notes).strip()


def build(spec_path: Path, out: Path, template: Path) -> Report:
    data = yaml.safe_load(spec_path.read_text(encoding="utf-8"))
    if not isinstance(data, dict) or "slides" not in data:
        raise SpecError("deck.yaml needs a top-level 'slides:' list (and optional 'meta:')")
    meta = data.get("meta") or {}
    ctx = Ctx(base=spec_path.parent, meta=meta, report=Report())
    prs = Presentation(str(template))
    if [lay.name for lay in prs.slide_layouts] != [lay.name for lay in D.LAYOUTS]:
        raise SpecError(f"{template} does not match design.py; rerun make_template.py")
    prs.core_properties.title = str(meta.get("title") or "")
    prs.core_properties.author = str(meta.get("author") or "")

    n = 0
    for i, spec in enumerate(data["slides"], 1):
        if not isinstance(spec, dict):
            raise SpecError(f"slide entry {i} is not a mapping")
        key = spec.get("layout")
        if key in CONTENT_LAYOUTS and not spec.get("backup"):
            ctx.content_slides += 1
            if not spec.get("notes"):
                ctx.slides_without_notes.append(str(n + 1))
        if spec.get("notes"):
            ctx.slides_with_notes.append(str(n + 1))
        # figure builds: image given as a list -> one slide per stage
        img = spec.get("image")
        if isinstance(img, list):
            # highlight boxes appear on the last stage only
            early = {k: v for k, v in spec.items() if k != "highlight"}
            for j, stage in enumerate(img):
                n += 1
                build_slide(prs, ctx, spec if j == len(img) - 1 else early, n,
                            images={"image": str(stage)})
            continue
        if spec.get("reveal"):
            role = next((r for r in ("text", "points") if r in spec), None)
            count = sum(1 for lvl, _ in flatten_items(spec.get(role) or []) if lvl == 0)
            for k in range(1, max(count, 1) + 1):
                n += 1
                build_slide(prs, ctx, spec, n, visible=k)
            continue
        n += 1
        build_slide(prs, ctx, spec, n)

    # deck-level checks
    r = ctx.report
    minutes = meta.get("duration_min")
    if minutes:
        hi = float(minutes) * 1.2
        if ctx.content_slides > hi:
            r.warn("deck", f"{ctx.content_slides} content slides for a {minutes}-minute talk "
                           f"(about one per minute; move extras to backup)")
    # Speaker notes are written only when Boss asks for them (meta.notes: true).
    if meta.get("notes") is True:
        if ctx.slides_without_notes:
            r.warn("deck", f"content slides without speaker notes: {', '.join(ctx.slides_without_notes)}")
    elif ctx.slides_with_notes:
        r.warn("deck", f"speaker notes on slide(s) {', '.join(ctx.slides_with_notes)} but nobody asked "
                       "for notes; remove them, or set meta.notes: true when Boss asks")
    for kind in sorted(MEASURE.missing):
        r.warn("deck", f"font for '{kind}' text not found; overflow checks are approximate")

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return r


def print_layouts() -> None:
    for lay in D.LAYOUTS:
        roles = [s for s in lay.slots if s.kind != "sldNum"]
        print(f"{lay.key:12s} ({lay.name})")
        for s in roles:
            size = f"{s.w:.2f} x {s.h:.2f} in"
            print(f"    {s.role:18s} {s.kind:8s} {size:18s} {s.prompt}")
        if lay.key == "table":
            print(f"    {'table':18s} {'table':8s} {D.CONTENT_W:.2f} x {D.BODY_H:.2f} in")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec", nargs="?", type=Path, help="deck.yaml")
    ap.add_argument("-o", "--out", type=Path, help="output .pptx (default: next to the spec)")
    ap.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    ap.add_argument("--layouts", action="store_true", help="list layouts and their fields")
    args = ap.parse_args()
    if args.layouts:
        print_layouts()
        return
    if args.spec is None:
        ap.error("give a deck.yaml or --layouts")
    out = args.out or args.spec.with_suffix(".pptx")
    try:
        report = build(args.spec, out, args.template)
    except SpecError as e:
        print(f"ERROR {e}", file=sys.stderr)
        sys.exit(2)
    for issue in report.issues:
        print(issue)
    n_err, n_warn = len(report.errors), len(report.issues) - len(report.errors)
    print(f"wrote {out}: {n_err} error(s), {n_warn} warning(s)")
    sys.exit(1 if n_err else 0)


if __name__ == "__main__":
    main()
