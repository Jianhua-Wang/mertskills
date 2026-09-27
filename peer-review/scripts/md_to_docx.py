#!/usr/bin/env python3
"""Convert a Markdown review report into a clean Word document for journal submission.

The Markdown subset is the one used by the report templates:
  # / ## / ### headings
  numbered items ("1." "12)" "(3)")        the literal number is kept, so it matches
                                           cross-references such as "see comment 7"
  bullets ("-", "*", "+"), nested by indentation
  paragraphs (consecutive lines are joined; a line ending in two spaces or a
              backslash forces a line break)
  **bold**, *italic* or _italic_, `code`, [text](url), ~~strikethrough~~
  "> quote" lines, pipe tables, and "---" horizontal rules
  HTML comments (<!-- ... -->) are dropped, so private notes can stay in the draft

Word styles are plain (single font, black headings), since journal systems paste the
text into web forms or send it on as-is.

The output never overwrites an existing file unless --overwrite is given: if
comments.docx exists, comments_v2.docx is written instead (then _v3, ...). The path
actually written is printed on stdout.

Examples:
  md_to_docx.py report.md                       # -> report.docx next to it
  md_to_docx.py report.md -o "/path/Review_comments.docx"
  md_to_docx.py 评审意见.md --font "Times New Roman" --east-asia-font SimSun --size 12
"""

from __future__ import annotations

import argparse
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Inches
except ImportError:
    sys.exit("python-docx is required: pip install python-docx")

HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
NUMBERED_RE = re.compile(r"^(\s*)(\(?\d+(?:\.\d+)*[.)]|\(\d+\)|[A-Za-z][.)])\s+(.*)$")
BULLET_RE = re.compile(r"^(\s*)[-*+•]\s+(.*)$")
RULE_RE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
QUOTE_RE = re.compile(r"^\s*>\s?(.*)$")
FENCE_RE = re.compile(r"^\s*(```|~~~)")
# A short label line such as "Minor:" or "**Figures and tables:**" ends any open list.
LABEL_RE = re.compile(r"^(\*\*|__)?[^\s*_|>#-][^:]{0,60}:(\*\*|__)?\s*$")

# Inline tokens, tried in order at each position.
INLINE_RE = re.compile(
    r"(?P<code>`[^`]+`)"
    r"|(?P<link>\[(?P<ltext>[^\]]+)\]\((?P<lurl>[^)\s]+)\))"
    r"|(?P<bolditalic>\*\*\*(?!\s)(?P<bi>.+?)(?<!\s)\*\*\*)"
    r"|(?P<bold>\*\*(?!\s)(?P<b>.+?)(?<!\s)\*\*|__(?!\s)(?P<b2>.+?)(?<!\s)__)"
    r"|(?P<strike>~~(?!\s)(?P<s>.+?)(?<!\s)~~)"
    r"|(?P<italic>(?<![\w*])\*(?!\s)(?P<i>.+?)(?<!\s)\*(?![\w*])|(?<![\w])_(?!\s)(?P<i2>.+?)(?<!\s)_(?![\w]))"
    r"|(?P<escape>\\[\\`*_{}\[\]()#+\-.!|~>])"
)


def set_run_font(run, font: str, east_asia: str | None, size: float | None = None):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), font)
    if east_asia:
        rfonts.set(qn("w:eastAsia"), east_asia)
    if size:
        run.font.size = Pt(size)


def setup_styles(doc, font: str, east_asia: str | None, size: float):
    for name in ("Normal", "List Bullet", "List Bullet 2", "List Bullet 3", "Quote", "Table Grid"):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = font
        if name == "Normal":
            st.font.size = Pt(size)
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), font)
        if east_asia:
            rfonts.set(qn("w:eastAsia"), east_asia)
    normal = doc.styles["Normal"].paragraph_format
    normal.space_after = Pt(6)
    normal.line_spacing = 1.15
    for level, pts in ((1, size + 4), (2, size + 2), (3, size + 1), (4, size), (5, size), (6, size)):
        try:
            st = doc.styles["Heading %d" % level]
        except KeyError:
            continue
        st.font.name = font
        st.font.size = Pt(pts)
        st.font.bold = True
        st.font.italic = level >= 4
        st.font.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), font)
        for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rfonts.get(qn(theme_attr)) is not None:
                del rfonts.attrib[qn(theme_attr)]
        if east_asia:
            rfonts.set(qn("w:eastAsia"), east_asia)
        pf = st.paragraph_format
        pf.space_before = Pt(12 if level <= 2 else 8)
        pf.space_after = Pt(4)
        pf.keep_with_next = True


def add_inline(par, text: str, font: str, east_asia: str | None, base=None):
    """Add Markdown-formatted inline text to a paragraph as runs."""
    base = dict(base or {})
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            _run(par, text[pos:m.start()], font, east_asia, **base)
        if m.group("code"):
            _run(par, m.group("code")[1:-1], "Courier New", None, **base)
        elif m.group("link"):
            label, url = m.group("ltext"), m.group("lurl")
            add_inline(par, label, font, east_asia, base)
            if url.rstrip("/") != label.rstrip("/"):
                _run(par, " (%s)" % url, font, east_asia, **base)
        elif m.group("bolditalic"):
            add_inline(par, m.group("bi"), font, east_asia, dict(base, bold=True, italic=True))
        elif m.group("bold"):
            add_inline(par, m.group("b") or m.group("b2"), font, east_asia, dict(base, bold=True))
        elif m.group("strike"):
            add_inline(par, m.group("s"), font, east_asia, dict(base, strike=True))
        elif m.group("italic"):
            add_inline(par, m.group("i") or m.group("i2"), font, east_asia, dict(base, italic=True))
        elif m.group("escape"):
            _run(par, m.group("escape")[1], font, east_asia, **base)
        pos = m.end()
    if pos < len(text):
        _run(par, text[pos:], font, east_asia, **base)


def _run(par, text, font, east_asia, bold=False, italic=False, strike=False):
    parts = text.split("\n")
    run = None
    for i, part in enumerate(parts):
        run = par.add_run(part)
        set_run_font(run, font, east_asia)
        run.bold = bold or None
        run.italic = italic or None
        if strike:
            run.font.strike = True
        if i < len(parts) - 1:
            run.add_break(WD_BREAK.LINE)
    return run


def join_lines(lines: list[str]) -> str:
    """Join soft-wrapped lines; keep hard breaks (two trailing spaces or a backslash)."""
    out = ""
    for i, ln in enumerate(lines):
        hard = ln.endswith("  ") or ln.endswith("\\")
        piece = ln.rstrip()
        if piece.endswith("\\"):
            piece = piece[:-1].rstrip()
        piece = piece.strip() if i else piece.lstrip()
        if not out:
            out = piece
        elif out.endswith("\n"):
            out += piece
        else:
            # No space between two CJK characters when joining wrapped Chinese text.
            if out and piece and _is_cjk(out[-1]) and _is_cjk(piece[0]):
                out += piece
            else:
                out += " " + piece
        if hard and i < len(lines) - 1:
            out += "\n"
    return out


def _is_cjk(ch: str) -> bool:
    code = ord(ch)
    return (0x4E00 <= code <= 0x9FFF or 0x3400 <= code <= 0x4DBF or 0x3000 <= code <= 0x303F
            or 0xFF00 <= code <= 0xFFEF)


def split_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|") and not s.endswith("\\|"):
        s = s[:-1]
    cells, cur, esc = [], "", False
    for ch in s:
        if esc:
            cur += ch
            esc = False
        elif ch == "\\":
            cur += ch
            esc = True
        elif ch == "|":
            cells.append(cur.strip())
            cur = ""
        else:
            cur += ch
    cells.append(cur.strip())
    return [c.replace("\\|", "|") for c in cells]


class ListLevels:
    """Nesting level from marker indentation, whatever the indent width (2, 3 or 4 spaces)."""

    def __init__(self):
        self.stack: list[int] = []

    def level(self, spaces: str) -> int:
        width = len(spaces.replace("\t", "    "))
        while self.stack and self.stack[-1] >= width:
            self.stack.pop()
        self.stack.append(width)
        return min(len(self.stack) - 1, 5)

    def reset(self):
        self.stack = []


def convert(md_text: str, font: str, east_asia: str | None, size: float, number_indent: float):
    doc = Document()
    setup_styles(doc, font, east_asia, size)
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(1)

    md_text = re.sub(r"<!--.*?-->", "", md_text, flags=re.S)
    lines = md_text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    i, n = 0, len(lines)
    para_buf: list[str] = []
    para_indent = None  # left indent (inches) when a paragraph continues a list item
    # The item currently open for continuation lines: (kind, paragraph, lines, label, text indent)
    item = None
    # Text indent of the last list item, so indented paragraphs after a blank line stay
    # inside that item (multi-paragraph comments). Cleared by any block at column 0.
    list_indent = None
    levels = ListLevels()

    def flush_para():
        nonlocal para_buf, para_indent
        if para_buf:
            p = doc.add_paragraph()
            if para_indent:
                p.paragraph_format.left_indent = Inches(para_indent)
            add_inline(p, join_lines(para_buf), font, east_asia)
            para_buf = []
            para_indent = None

    def flush_item():
        nonlocal item, list_indent
        if item is not None:
            kind, p, buf, label, text_indent = item
            if label:
                r = p.add_run(label + "\t")
                set_run_font(r, font, east_asia)
            add_inline(p, join_lines(buf), font, east_asia)
            list_indent = text_indent
            item = None

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if FENCE_RE.match(line):
            flush_para(); flush_item()
            list_indent = None
            levels.reset()
            fence = FENCE_RE.match(line).group(1)
            i += 1
            code_lines = []
            while i < n and not lines[i].strip().startswith(fence):
                code_lines.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(code_lines))
            set_run_font(r, "Courier New", None, size - 1)
            continue

        if not stripped:
            flush_para(); flush_item()
            i += 1
            continue

        m = HEADING_RE.match(line)
        if m:
            flush_para(); flush_item()
            list_indent = None
            levels.reset()
            level = len(m.group(1))
            h = doc.add_heading(level=min(level, 6))
            add_inline(h, m.group(2), font, east_asia)
            i += 1
            continue

        if RULE_RE.match(line) and not para_buf:
            flush_item()
            list_indent = None
            levels.reset()
            p = doc.add_paragraph()
            ppr = p._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "999999")):
                bottom.set(qn(k), v)
            bdr.append(bottom)
            ppr.append(bdr)
            i += 1
            continue

        if "|" in line and i + 1 < n and TABLE_SEP_RE.match(lines[i + 1]):
            flush_para(); flush_item()
            list_indent = None
            levels.reset()
            header = split_row(line)
            rows = []
            i += 2
            while i < n and lines[i].strip() and "|" in lines[i]:
                rows.append(split_row(lines[i]))
                i += 1
            ncols = max([len(header)] + [len(r) for r in rows])
            table = doc.add_table(rows=1 + len(rows), cols=ncols)
            try:
                table.style = doc.styles["Table Grid"]
            except KeyError:
                pass
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_idx, row in enumerate([header] + rows):
                for c_idx in range(ncols):
                    cell = table.cell(r_idx, c_idx)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(0)
                    txt = row[c_idx] if c_idx < len(row) else ""
                    add_inline(p, txt.replace("<br>", "\n").replace("<br/>", "\n"), font, east_asia,
                               {"bold": r_idx == 0})
            doc.add_paragraph()
            continue

        m = QUOTE_RE.match(line)
        if m:
            flush_para(); flush_item()
            indent = list_indent if (list_indent and line[:1] in (" ", "\t")) else 0
            if not indent:
                list_indent = None
                levels.reset()
            quote = []
            while i < n and QUOTE_RE.match(lines[i]):
                quote.append(QUOTE_RE.match(lines[i]).group(1))
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(indent + 0.4)
            add_inline(p, join_lines(quote), font, east_asia, {"italic": True})
            continue

        m = NUMBERED_RE.match(line)
        if m:
            # A numbered line always starts an item, even right after "Major comments:".
            flush_para(); flush_item()
            level = levels.level(m.group(1))
            text_indent = number_indent * (level + 1)
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Inches(text_indent)
            pf.first_line_indent = Inches(-number_indent)
            pf.tab_stops.add_tab_stop(Inches(text_indent))
            item = ("num", p, [m.group(3)], m.group(2), text_indent)
            i += 1
            continue

        m = BULLET_RE.match(line)
        if m and not RULE_RE.match(line):
            flush_para(); flush_item()
            level = levels.level(m.group(1))
            # A hanging bullet: the glyph sits under the parent item's text.
            text_indent = number_indent * level + 0.25
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Inches(text_indent)
            pf.first_line_indent = Inches(-0.25)
            pf.tab_stops.add_tab_stop(Inches(text_indent))
            item = ("bullet", p, [m.group(2)], "•◦▪"[level % 3], text_indent)
            i += 1
            continue

        if item is not None and LABEL_RE.match(line):
            flush_item()
            list_indent = None
            levels.reset()

        # Continuation of an open list item (lazy continuation), an indented paragraph
        # that belongs to the previous list item, or an ordinary paragraph.
        if item is not None:
            item[2].append(line)
        else:
            if not para_buf:
                if list_indent and line[:1] in (" ", "\t"):
                    para_indent = list_indent
                else:
                    list_indent = None
                    levels.reset()
            para_buf.append(line)
        i += 1

    flush_para(); flush_item()
    return doc


def next_free(path: Path) -> Path:
    if not path.exists():
        return path
    stem = re.sub(r"_v\d+$", "", path.stem)
    k = 2
    while True:
        cand = path.with_name("%s_v%d%s" % (stem, k, path.suffix))
        if not cand.exists():
            return cand
        k += 1


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("markdown", help="the Markdown report")
    ap.add_argument("-o", "--output", help="output .docx (default: next to the Markdown file)")
    ap.add_argument("--font", default="Times New Roman", help="Latin font (default Times New Roman)")
    ap.add_argument("--east-asia-font", default=None,
                    help="font for Chinese/Japanese/Korean text, e.g. SimSun or 宋体 "
                         "(default: SimSun when the text contains CJK characters)")
    ap.add_argument("--size", type=float, default=12, help="body font size in points (default 12)")
    ap.add_argument("--number-indent", type=float, default=0.3,
                    help="hanging indent for numbered items, in inches (default 0.3)")
    ap.add_argument("--overwrite", action="store_true", help="replace the output file if it exists")
    args = ap.parse_args(argv)

    src = Path(args.markdown).expanduser()
    text = src.read_text(encoding="utf-8-sig")
    east_asia = args.east_asia_font
    if east_asia is None and any(_is_cjk(ch) for ch in text):
        east_asia = "SimSun"
    doc = convert(text, args.font, east_asia, args.size, args.number_indent)

    out = Path(args.output).expanduser() if args.output else src.with_suffix(".docx")
    if out.suffix.lower() != ".docx":
        out = out.with_suffix(".docx")
    out.parent.mkdir(parents=True, exist_ok=True)
    if not args.overwrite:
        target = next_free(out)
        if target != out:
            print("[note] %s exists; writing %s instead (use --overwrite to replace)" % (out.name, target.name),
                  file=sys.stderr)
        out = target
    # Blank identifying metadata: reports are often forwarded to authors as files.
    now = datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None)
    props = doc.core_properties
    props.author = props.last_modified_by = props.title = props.comments = ""
    props.keywords = props.subject = props.category = ""
    props.created = props.modified = now
    props.revision = 1
    doc.save(str(out))
    print(out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
